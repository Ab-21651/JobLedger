import os

import openpyxl
import pytest

from jobledger.db import queries as q
from jobledger.tools import excel_exporter as xl


def _seed(ws, title="Backend Python Developer", company="Acme Cloud"):
    jid, _ = q.upsert_job(ws.conn, ws.pid, None, title, company, location="London", arrangement="hybrid",
                          employment_type="fulltime", apply_url="https://careers.acme.example/jobs/101")
    return jid


def test_safe_cell_prefixes_formula_characters():
    for bad in ("=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)"):
        assert xl.safe_cell(bad).startswith("'")
    assert xl.safe_cell("Acme") == "Acme" and xl.safe_cell(5) == 5


def test_export_writes_sheets_and_no_live_formulas(ws):
    _seed(ws, title="=HYPERLINK(\"http://evil.example\",\"x\")", company="+Evil")
    res = xl.export_excel(ws.conn, ws.pid, ws.settings.export_dir)
    wb = openpyxl.load_workbook(res["path"])
    assert wb.sheetnames[:3] == ["Register", "Conflicts", "About"]
    for row in wb["Register"].iter_rows(min_row=2):
        for cell in row:
            assert cell.data_type != "f", cell.coordinate
    wb.close()


def test_repeat_export_is_stable_and_row_count_matches_db(ws):
    _seed(ws)
    r1 = xl.export_excel(ws.conn, ws.pid, ws.settings.export_dir)
    r2 = xl.export_excel(ws.conn, ws.pid, ws.settings.export_dir)
    assert r1["path"] == r2["path"] and r2["rows"] == 1 and not r2["locked_fallback"]


def test_locked_file_falls_back_to_new_file_and_original_untouched(ws, monkeypatch):
    _seed(ws)
    first = xl.export_excel(ws.conn, ws.pid, ws.settings.export_dir)
    before = os.path.getsize(first["path"])
    real = os.replace

    def locked(src, dst, *a, **k):
        if os.path.abspath(dst) == os.path.abspath(first["path"]):
            raise PermissionError("file is open in Excel")
        return real(src, dst, *a, **k)
    monkeypatch.setattr(os, "replace", locked)
    res = xl.export_excel(ws.conn, ws.pid, ws.settings.export_dir)
    assert res["locked_fallback"] and res["path"] != first["path"] and os.path.exists(res["path"])
    assert os.path.getsize(first["path"]) == before


def test_each_profile_gets_its_own_file_and_rows(ws):
    _seed(ws)
    rid = ws.riley()
    a = xl.export_excel(ws.conn, ws.pid, ws.settings.export_dir, filename=f"jobledger_register_p{ws.pid}.xlsx")
    b = xl.export_excel(ws.conn, rid, ws.settings.export_dir, filename=f"jobledger_register_p{rid}.xlsx")
    assert a["path"] != b["path"] and a["rows"] == 1 and b["rows"] == 0


def test_unknown_profile_refused(ws):
    with pytest.raises(q.ProfileIsolationError):
        xl.export_excel(ws.conn, 999, ws.settings.export_dir)
