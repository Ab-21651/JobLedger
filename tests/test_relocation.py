"""Willing to relocate (profiles.open_to_relocation): 'no' excludes confirmed far on-site/hybrid work; NULL and 'yes'
behave exactly as before. Offline: deterministic scoring, distances pre-seeded in place_distances."""
import dataclasses
import itertools

import pytest

from jobledger import cli
from jobledger.config import reset_settings
from jobledger.db import queries as q
from jobledger.db import schema
from jobledger.db.queries import norm_city
from jobledger.tools import matcher as m
from jobledger.tools.excel_exporter import export_excel
from jobledger.tools.profile_reader import read_full_profile

CV_TEXT = "Python developer, five years building web services."
JD = "Requirements: Python. Full-time."          # no arrangement wording: the job row decides arrangement
_n = itertools.count(1)


def _profile(conn, email, relocation=None, location="Islamabad"):
    return q.create_profile(conn, "Ali Reloc", email, location=location, open_to_relocation=relocation)


def _cv(conn, pid):
    q.add_cv_version(conn, pid, "cv-reloc", CV_TEXT, "Python", 5, "BSc")


def _seed(conn, a, b, km):
    q.save_place_distance(conn, norm_city(a), norm_city(b), float(km), "llm")


def _job(conn, pid, location, arrangement, title="Backend Engineer"):
    n = next(_n)
    url = f"https://reloc{n}.example/job"
    jid, _ = q.upsert_job(conn, pid, None, title, f"Reloc Co {n}", location=location, arrangement=arrangement,
                          employment_type="fulltime", apply_url=url)
    sid = q.save_jd_snapshot(conn, jid, url, url, "ok", content_text=JD, verification_status="verified")
    return jid, conn.execute("SELECT * FROM jd_snapshots WHERE id=?", (sid,)).fetchone()


def _eval(conn, pid, jid, snap, llm=None):
    return m.evaluate_job(conn, read_full_profile(conn, pid), q.get_job(conn, jid), snap, llm=llm)


# ---------------------------------------------------------------- default behaviour is unchanged

GRID = [("Berlin", "onsite"), ("Berlin", "hybrid"), ("Rawalpindi", "onsite"), ("Lahore", "hybrid"),
        ("London (Hybrid)", "hybrid"), ("Remote - Berlin", None), ("Berlin", None), ("Islamabad", "onsite"),
        (None, "onsite"), ("Karachi", "remote"), ("Berlin", "fulltime")]


def test_unset_and_yes_score_exactly_the_same(ws):
    for a, b, km in [("Islamabad", "Berlin", 5000), ("Islamabad", "Rawalpindi", 14), ("Islamabad", "Lahore", 370)]:
        _seed(ws.conn, a, b, km)
    unset, yes = _profile(ws.conn, "grid.unset@example.com"), _profile(ws.conn, "grid.yes@example.com", "yes")
    results = {}
    for name, pid in (("unset", unset), ("yes", yes)):
        _cv(ws.conn, pid)
        prof = read_full_profile(ws.conn, pid)
        assert prof.open_to_relocation == "yes" and "relocation" not in m.hard_filters(prof, q.get_job(
            ws.conn, _job(ws.conn, pid, "Berlin", "onsite")[0]), JD)
        results[name] = []
        for loc, arr in GRID:
            jid, snap = _job(ws.conn, pid, loc, arr)
            o = _eval(ws.conn, pid, jid, snap)
            results[name].append((o.status, o.total_score, o.hard_filter_pass, o.components, o.rationale))
    assert results["unset"] == results["yes"]
    assert not any(r[0] == "excluded" for r in results["yes"])


# ---------------------------------------------------------------- 'no' and the relocation rule

def test_no_excludes_a_far_hybrid_job_and_says_why(ws):
    _seed(ws.conn, "Islamabad", "Berlin", 5000)
    pid = _profile(ws.conn, "ali.hybrid@example.com", "no")
    _cv(ws.conn, pid)
    jid, snap = _job(ws.conn, pid, "Berlin", "hybrid")
    out = _eval(ws.conn, pid, jid, snap)
    assert out.status == "excluded" and out.hard_filter_pass is False
    assert "requires relocation" in out.rationale and "job in Berlin is ~5000 km away" in out.rationale
    assert "relocation" in q.get_latest_evaluation(ws.conn, jid, pid)["hard_filter_details"]


def test_no_keeps_an_onsite_job_in_rawalpindi_at_14_km(ws):
    _seed(ws.conn, "Islamabad", "Rawalpindi", 14)
    pid = _profile(ws.conn, "ali.near@example.com", "no")
    _cv(ws.conn, pid)
    jid, snap = _job(ws.conn, pid, "Rawalpindi", "onsite")
    out = _eval(ws.conn, pid, jid, snap)
    assert out.status != "excluded" and out.components["location"] == 1.0
    assert "within commuting distance (~14 km" in out.rationale


def test_no_keeps_a_remote_job_however_far(ws):
    _seed(ws.conn, "Islamabad", "Berlin", 5000)
    pid = _profile(ws.conn, "ali.remote@example.com", "no")
    _cv(ws.conn, pid)
    jid, snap = _job(ws.conn, pid, "Berlin", "remote")
    jid2, snap2 = _job(ws.conn, pid, "Remote, Germany", None)
    for j, s in ((jid, snap), (jid2, snap2)):
        assert _eval(ws.conn, pid, j, s).status != "excluded"
    prof = read_full_profile(ws.conn, pid)
    assert m.hard_filters(prof, q.get_job(ws.conn, jid), JD, distance_km=5000)["relocation"]["detail"] == "remote job"


def test_no_with_arrangement_unknown_far_is_flagged_not_excluded(ws):
    _seed(ws.conn, "Islamabad", "Berlin", 5000)
    pid = _profile(ws.conn, "ali.unknown.arr@example.com", "no")
    _cv(ws.conn, pid)
    jid, snap = _job(ws.conn, pid, "Berlin", None)
    out = _eval(ws.conn, pid, jid, snap)
    assert out.status != "excluded" and out.status in ("scored", "provisional")
    assert "relocation_possibly_needed" in out.flags
    prof = read_full_profile(ws.conn, pid)
    rule = m.hard_filters(prof, q.get_job(ws.conn, jid), JD, distance_km=5000)["relocation"]
    assert rule["result"] == "unknown" and "may need relocation" in rule["detail"]


def test_no_with_distance_unknown_is_not_excluded(ws):
    pid = _profile(ws.conn, "ali.unknown.dist@example.com", "no")
    _cv(ws.conn, pid)
    jid, snap = _job(ws.conn, pid, "Berlin", "onsite")
    out = _eval(ws.conn, pid, jid, snap, llm=None)                  # no cache, no LLM: distance not known
    assert out.status != "excluded" and "relocation_possibly_needed" not in out.flags
    rule = m.hard_filters(read_full_profile(ws.conn, pid), q.get_job(ws.conn, jid), JD)["relocation"]
    assert rule["result"] == "unknown" and rule["detail"] == "distance to Berlin not known"


def test_no_in_the_same_city_passes(ws):
    pid = _profile(ws.conn, "ali.same@example.com", "no", location="Berlin")
    _cv(ws.conn, pid)
    jid, snap = _job(ws.conn, pid, "Berlin", "onsite")
    assert _eval(ws.conn, pid, jid, snap).status != "excluded"
    rule = m.hard_filters(read_full_profile(ws.conn, pid), q.get_job(ws.conn, jid), JD, distance_km=None)["relocation"]
    assert rule["result"] == "pass"


# ---------------------------------------------------------------- switching back re-scores

def test_switching_back_to_yes_rescores_the_job(ws):
    _seed(ws.conn, "Islamabad", "Berlin", 5000)
    pid = _profile(ws.conn, "ali.switch@example.com", "yes")
    _cv(ws.conn, pid)
    jid, snap = _job(ws.conn, pid, "Berlin", "hybrid")
    first = _eval(ws.conn, pid, jid, snap)                          # yes: scored, not excluded
    assert first.status != "excluded"
    q.update_profile_fields(ws.conn, pid, {"open_to_relocation": "no"}, "not moving")
    assert _eval(ws.conn, pid, jid, snap).status == "excluded"
    q.update_profile_fields(ws.conn, pid, {"open_to_relocation": "yes"}, "changed my mind")
    prof = read_full_profile(ws.conn, pid)
    latest = q.get_latest_evaluation(ws.conn, jid, pid)
    assert latest["status"] == "excluded" and latest["prefs_version"] != prof.prefs_version   # stale: agent re-scores
    again = _eval(ws.conn, pid, jid, snap)
    assert again.inserted is True and again.status == first.status and again.total_score == first.total_score
    latest = q.get_latest_evaluation(ws.conn, jid, pid)
    assert latest["prefs_version"] == prof.prefs_version and latest["status"] == first.status
    assert "requires relocation" not in latest["match_rationale"]


# ---------------------------------------------------------------- validation, storage, history

@pytest.mark.parametrize("bad", ["maybe", "sometimes", "1"])
def test_invalid_value_is_refused_by_the_queries(ws, bad):
    pid = _profile(ws.conn, f"ali.bad{bad}@example.com")
    with pytest.raises(ValueError, match="'yes' or 'no'"):
        q.update_profile_fields(ws.conn, pid, {"open_to_relocation": bad}, "x")
    with pytest.raises(ValueError, match="'yes' or 'no'"):
        q.create_profile(ws.conn, "Bad", f"bad{bad}.create@example.com", open_to_relocation=bad)
    assert q.get_profile_by_email(ws.conn, f"bad{bad}.create@example.com") is None
    assert ws.conn.execute("SELECT COUNT(*) FROM preferences_history WHERE profile_id=? AND "
                           "field_name='open_to_relocation'", (pid,)).fetchone()[0] == 0


def test_values_are_case_insensitive_and_stored_lowercase(ws):
    pid = _profile(ws.conn, "ali.case@example.com")
    q.update_profile_fields(ws.conn, pid, {"open_to_relocation": "NO"}, "x")
    assert ws.conn.execute("SELECT open_to_relocation FROM profiles WHERE id=?", (pid,)).fetchone()[0] == "no"
    assert read_full_profile(ws.conn, pid).open_to_relocation == "no"


def test_update_writes_history_and_bumps_prefs_version(ws):
    pid = _profile(ws.conn, "ali.history@example.com")
    before = q.prefs_version(ws.conn, pid)
    applied = q.update_profile_fields(ws.conn, pid, {"open_to_relocation": "no"}, "not moving")
    assert applied == [{"field": "open_to_relocation", "old": None, "new": "no"}]
    assert q.prefs_version(ws.conn, pid) > before
    row = ws.conn.execute("SELECT old_value, new_value, reason FROM preferences_history WHERE field_name="
                          "'open_to_relocation'").fetchone()
    assert tuple(row) == (None, "no", "not moving")


def test_preference_snapshot_and_input_hash_follow_the_setting(ws):
    pid = _profile(ws.conn, "ali.hash@example.com")
    _cv(ws.conn, pid)
    jid, snap = _job(ws.conn, pid, "Berlin", "onsite")
    prof, job = read_full_profile(ws.conn, pid), q.get_job(ws.conn, jid)
    no = dataclasses.replace(prof, open_to_relocation="no")
    assert prof.preference_snapshot()["open_to_relocation"] == "yes"
    assert no.preference_snapshot() != prof.preference_snapshot()
    assert m.compute_input_hash(prof, job, snap, m.WEIGHTS, False) != m.compute_input_hash(no, job, snap, m.WEIGHTS, False)


# ---------------------------------------------------------------- migration 6

def test_migration_6_upgrades_a_v5_database_in_place(tmp_path, monkeypatch):
    path = str(tmp_path / "v5.db")
    full = schema.MIGRATIONS
    monkeypatch.setattr(schema, "MIGRATIONS", full[:5])
    conn = schema.init_db(path)
    assert schema.current_version(conn) == 5
    pid = q.create_profile(conn, "Old Five", "old.five@example.com", location="Islamabad")
    conn.close()
    monkeypatch.setattr(schema, "MIGRATIONS", full)
    conn = schema.init_db(path)
    assert schema.current_version(conn) == 6 == schema.SCHEMA_VERSION
    assert "open_to_relocation" in {r["name"] for r in conn.execute("PRAGMA table_info(profiles)")}
    assert conn.execute("SELECT open_to_relocation FROM profiles WHERE id=?", (pid,)).fetchone()[0] is None
    assert read_full_profile(conn, pid).open_to_relocation == "yes"
    assert q.get_profile(conn, pid)["name"] == "Old Five"
    conn.close()


# ---------------------------------------------------------------- the register

def test_excluded_job_appears_in_the_register_with_its_reason(ws, tmp_path):
    from openpyxl import load_workbook
    _seed(ws.conn, "Islamabad", "Berlin", 5000)
    pid = _profile(ws.conn, "ali.register@example.com", "no")
    _cv(ws.conn, pid)
    jid, snap = _job(ws.conn, pid, "Berlin", "hybrid", title="Far Away Engineer")
    assert _eval(ws.conn, pid, jid, snap).status == "excluded"
    res = export_excel(ws.conn, pid, export_dir=str(tmp_path), filename="reg.xlsx")
    rows = [[str(c.value) for c in row if c.value is not None] for row in load_workbook(res["path"]).active.iter_rows()]
    hit = [r for r in rows if "Far Away Engineer" in r]
    assert hit and "EXCLUDED" in hit[0] and any("requires relocation" in v for v in hit[0])


# ---------------------------------------------------------------- CLI

@pytest.fixture
def cli_env(tmp_path, monkeypatch):
    for k in ("IMAP_USER", "IMAP_PASSWORD", "LLM_PROVIDER", "LLM_API_KEY"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("DB_PATH", str(tmp_path / "cli.db"))
    monkeypatch.setenv("EXPORT_DIR", str(tmp_path / "exports"))
    monkeypatch.setenv("LOG_DIR", str(tmp_path / "logs"))
    monkeypatch.chdir(tmp_path)
    reset_settings()
    yield tmp_path
    reset_settings()


def _run(capsys, *argv):
    code = cli.main(list(argv))
    return code, capsys.readouterr().out


def test_cli_profile_create_relocation_no_stores_no(cli_env, capsys):
    code, _ = _run(capsys, "profile", "create", "--name", "Ali", "--email", "ali.cli@example.com", "--relocation", "no")
    assert code == 0
    conn = schema.init_db(str(cli_env / "cli.db"))
    assert q.get_profile_by_email(conn, "ali.cli@example.com")["open_to_relocation"] == "no"
    conn.close()


def test_cli_profile_create_rejects_an_unknown_choice(cli_env, capsys):
    with pytest.raises(SystemExit) as exc:
        cli.main(["profile", "create", "--name", "Ali", "--email", "ali.badchoice@example.com", "--relocation", "maybe"])
    assert exc.value.code == 2 and "invalid choice" in capsys.readouterr().err


def test_cli_prefs_set_writes_history_and_reports_a_bad_value(cli_env, capsys):
    assert _run(capsys, "profile", "create", "--name", "Ali", "--email", "ali.prefs@example.com")[0] == 0
    code, out = _run(capsys, "prefs", "set", "--set", "open_to_relocation=maybe")
    assert code == 1 and out.startswith("error:") and "'yes' or 'no'" in out
    code, out = _run(capsys, "prefs", "set", "--set", "open_to_relocation=no", "--reason", "not moving")
    assert code == 0 and "open_to_relocation" in out
    conn = schema.init_db(str(cli_env / "cli.db"))
    row = conn.execute("SELECT old_value, new_value, reason FROM preferences_history WHERE field_name="
                       "'open_to_relocation'").fetchall()
    assert [tuple(r) for r in row] == [(None, "no", "not moving")]
    conn.close()
