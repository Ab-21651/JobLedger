"""
Tool H: Excel Register
Generates ranked Excel workbook from SQLite.
Source of truth is SQLite only.
Writes to temp file then replaces destination safely.
"""
from __future__ import annotations

import os
import shutil
import sqlite3
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from jobledger.logging_setup import get_logger

logger = get_logger("excel_exporter")

# Colours
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GREY_FONT = Font(color="808080", size=9)
LINK_FONT = Font(color="0563C1", underline="single", size=9)

COLUMNS = [
    ("Rank", 6),
    ("Job Title", 30),
    ("Company", 20),
    ("Location", 18),
    ("Arrangement", 12),
    ("Type", 10),
    ("Score %", 9),
    ("Skill", 8),
    ("Exp.", 8),
    ("Arr.", 8),
    ("Match Rationale", 40),
    ("Missing Info", 20),
    ("JD Status", 14),
    ("Stage", 12),
    ("Applied Date", 13),
    ("Last Updated", 13),
    ("Apply Link", 35),
    ("Next Action", 30),
]


def _score_fill(score: float) -> Optional[PatternFill]:
    if score >= 0.70:
        return GREEN_FILL
    if score >= 0.40:
        return YELLOW_FILL
    return RED_FILL


def _query_jobs_for_export(
    conn: sqlite3.Connection,
    profile_id: int,
    include_stages: Optional[list] = None,
) -> list[dict]:
    """Query jobs with evaluations and application stages."""
    stage_filter = ""
    params = [profile_id]

    if include_stages:
        placeholders = ",".join("?" * len(include_stages))
        stage_filter = f"AND COALESCE(a.stage, 'saved') IN ({placeholders})"
        params.extend(include_stages)

    sql = f"""
    SELECT
        j.id AS job_id,
        j.title,
        j.company,
        j.location,
        j.arrangement,
        j.employment_type,
        j.apply_url,
        j.jd_status,
        j.updated_at,
        e.total_score,
        e.skill_score,
        e.experience_score,
        e.arrangement_score,
        e.match_rationale,
        e.missing_info,
        COALESCE(a.stage, 'saved') AS stage,
        a.applied_at,
        a.updated_at AS app_updated_at
    FROM jobs j
    LEFT JOIN evaluations e ON e.job_id = j.id AND e.profile_id = j.profile_id
    LEFT JOIN applications a ON a.job_id = j.id
    WHERE j.profile_id = ?
    {stage_filter}
    ORDER BY COALESCE(e.total_score, 0) DESC, j.created_at DESC
    """
    rows = conn.execute(sql, params).fetchall()
    return [dict(r) for r in rows]


def export_excel(
    conn: sqlite3.Connection,
    profile_id: int,
    output_path: str,
    include_stages: Optional[list] = None,
) -> tuple[bool, str]:
    """
    Export ranked job register to Excel.
    Returns (success, message).
    """
    try:
        jobs = _query_jobs_for_export(conn, profile_id, include_stages)
    except Exception as e:
        logger.error(f"DB query failed for Excel export: {e}")
        return False, f"Database error: {e}"

    if not jobs:
        return False, "No jobs found for this profile to export"

    # Build workbook in memory / temp file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Job Register"

    # Header row
    for col_idx, (col_name, col_width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width
    ws.row_dimensions[1].height = 20

    # Freeze header
    ws.freeze_panes = "A2"

    # Data rows
    for rank, job in enumerate(jobs, start=1):
        row = rank + 1
        score = job.get("total_score") or 0.0

        from jobledger.tools.stage_manager import get_recommended_next_action
        stage = job.get("stage", "saved")
        applied_at = job.get("applied_at")
        days_applied = None
        if applied_at:
            try:
                dt = datetime.fromisoformat(applied_at)
                days_applied = (datetime.utcnow() - dt).days
            except Exception:
                pass

        next_action = get_recommended_next_action(stage, days_applied)

        rationale = (job.get("match_rationale") or "")[:200]
        missing = (job.get("missing_info") or "")[:100]
        apply_url = job.get("apply_url") or ""

        row_data = [
            rank,
            job.get("title") or "",
            job.get("company") or "",
            job.get("location") or "",
            job.get("arrangement") or "",
            job.get("employment_type") or "",
            f"{score * 100:.0f}%" if job.get("total_score") is not None else "N/A",
            f"{(job.get('skill_score') or 0) * 100:.0f}%",
            f"{(job.get('experience_score') or 0) * 100:.0f}%",
            f"{(job.get('arrangement_score') or 0) * 100:.0f}%",
            rationale,
            missing,
            job.get("jd_status") or "PENDING",
            stage,
            (applied_at or "")[:10],
            (job.get("app_updated_at") or job.get("updated_at") or "")[:10],
            apply_url,
            next_action,
        ]

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(size=9)

        # Colour-code by score
        fill = _score_fill(score)
        if fill:
            for col_idx in range(1, len(COLUMNS) + 1):
                ws.cell(row=row, column=col_idx).fill = fill

        # Hyperlink for apply URL
        if apply_url and apply_url.startswith("http"):
            url_cell = ws.cell(row=row, column=17)
            url_cell.hyperlink = apply_url
            url_cell.value = "Apply"
            url_cell.font = LINK_FONT

    # Write to temp file and replace destination
    try:
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
        os.close(tmp_fd)

        wb.save(tmp_path)

        # Replace destination
        shutil.move(tmp_path, output_path)
        logger.info(f"Excel exported: {output_path} | {len(jobs)} rows")
        return True, f"Exported {len(jobs)} jobs to {output_path}"

    except PermissionError:
        msg = f"File is locked (open in Excel?): {output_path}"
        logger.error(msg)
        try:
            os.unlink(tmp_path)
        except Exception:
            pass
        return False, msg
    except Exception as e:
        logger.error(f"Excel write error: {e}")
        try:
            os.unlink(tmp_path)
        except Exception:
            pass
        return False, f"Export error: {e}"
